"""公开来源和格式准备。所有路径均由本次运行目录生成，不读取历史项目。"""
import csv
import gzip
import io
import re
import shutil
import zipfile
from pathlib import Path
import numpy as np
import pandas as pd
from .io import download, dump, require, sha

GEO = 'https://ftp.ncbi.nlm.nih.gov/geo/'
MOUSE = GEO + 'series/GSE233nnn/GSE233208/suppl/'
MOUSE_TABLES = ['GSE233208_pilot_samples.csv.gz', 'GSE233208_5XFAD_8samples.csv.gz'] + [
    f'GSE233208_visium_5xFAD_{date}_sample_meta.csv.gz'
    for date in ('11-15-21', '11-24-21', '12-13-21', '12-20-21')]
URLS = {
    'seaad_donors.xlsx': 'https://cdn.prod.website-files.com/689cfbd308fa7373b604d290/68debdfdd1b8e9f8fd64dab0_sea-ad_cohort_donor_metadata_072524.xlsx',
    'mouse_tables.zip': 'https://static-content.springer.com/esm/art%3A10.1038%2Fs41588-024-01961-x/MediaObjects/41588_2024_1961_MOESM4_ESM.zip',
    'GSE118553_series_matrix.txt.gz': GEO + 'series/GSE118nnn/GSE118553/matrix/GSE118553_series_matrix.txt.gz',
    'GPL10558.annot.gz': GEO + 'platforms/GPL10nnn/GPL10558/annot/GPL10558.annot.gz',
    'orthologs.tsv': 'https://www.informatics.jax.org/downloads/reports/HOM_MouseHumanSequence.rpt',
    'mouse.rds.gz': MOUSE + 'GSE233208_5XFAD_seurat_processed_annotated.rds.gz',
    **{name: MOUSE + name for name in MOUSE_TABLES},
}


def fetch_small(raw):
    """下载元数据、基因注释和芯片矩阵；大计数对象另行处理。"""
    for name, url in URLS.items():
        if name != 'mouse.rds.gz':
            download(url, raw / name)


def scalar_id(value):
    """将整数样本编号规范成文本，不从含字母的动物编号中猜数字。"""
    text = str(value).strip()
    if re.fullmatch(r'\d+\.0', text):
        text = text[:-2]
    require(re.fullmatch(r'\d+', text) is not None, f'样本编号不是整数：{value}')
    return text


def cage_id(value):
    """保留作者的笼位分类标签；6.5与6是不同笼位，不能取整合并。

    Excel数值中的整数尾部统一去掉，文本标签只去首尾空格。
    缺失、无穷值和布尔值均不作为笼号使用。
    """
    require(not isinstance(value, (bool, np.bool_)) and not pd.isna(value),
            f'Cage.No笼位标签缺失或无效：{value}')
    if isinstance(value, (int, float, np.integer, np.floating)):
        require(np.isfinite(value), f'Cage.No笼位标签不是有限值：{value}')
        return str(int(value)) if value == int(value) else str(value)
    text = str(value).strip()
    require(text and text.lower() not in {'na', 'nan', 'none', 'null', 'inf', '-inf'},
            f'Cage.No笼位标签缺失或无效：{value}')
    return text


def age_month(value):
    """解析作者以月为单位的年龄；不将出生日期或取材日期当作年龄。"""
    text = str(value).strip().lower()
    match = re.fullmatch(r'(\d+(?:\.\d+)?)\s*(?:mo|months?|m)?', text)
    return float(match[1]) if match else np.nan


def genotype(value):
    """统一公开表格中的基因型缩写，未识别值保留缺失。"""
    text = str(value).upper().replace(' ', '').replace('-', '')
    return {'WT': 'WT', '5X': '5X', '5XFAD': '5X'}.get(text)


def mouse_manifest(raw, prepared):
    """用GEO表确定纳入资格，再以发表的S2表确定动物和笼位。

    同一SAMPLE只能对应一个动物；笼号和耳号组合也须唯一。重复切片不会被
    当作新动物。GEO缺失性别的试验样本不由其他表静默补入。
    """
    rows = []
    for name in MOUSE_TABLES:
        frame = pd.read_csv(raw / name)
        frame.columns = [str(c).strip().lstrip('\ufeff') for c in frame]
        require('SAMPLE' in frame, f'{name}缺少SAMPLE')
        frame = frame.loc[frame.SAMPLE.notna()].copy()
        frame['SAMPLE'] = frame.SAMPLE.map(scalar_id)
        frame['source_file'] = name
        rows.append(frame)
    source = pd.concat(rows, ignore_index=True)
    require(not source.SAMPLE.duplicated().any(), 'GEO样本表之间存在重复编号，需要人工核查来源')
    require({'Sex', 'Age', 'Condition'} <= set(source), 'GEO性别/年龄/条件列缺失')
    source['sex'] = source.Sex.astype(str).str.upper().str.strip()
    source['age'] = source.Age.map(age_month)
    source['genotype'] = source.Condition.map(genotype)
    source['eligible'] = source.sex.isin(['M', 'F']) & source.age.notna() & source.genotype.notna()
    source.to_csv(prepared / 'mouse_source_eligibility.csv', index=False)
    with zipfile.ZipFile(raw / 'mouse_tables.zip') as archive:
        names = [n for n in archive.namelist() if n.endswith('.xlsx') and not n.startswith('__MACOSX/')]
        require(len(names) == 1, f'发表补充资料工作簿不唯一：{names}')
        workbook = raw / 'published_mouse_tables.xlsx'
        with archive.open(names[0]) as src, workbook.open('wb') as dst:
            shutil.copyfileobj(src, dst, 4 * 1024**2)
    import openpyxl
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    require('Supp. Table 2' in book.sheetnames, '发表表格中未找到Supp. Table 2')
    iterator = book['Supp. Table 2'].iter_rows(values_only=True)
    header = None
    for _ in range(20):
        vals = next(iterator)
        clean = [str(v).strip() if v is not None else '' for v in vals]
        if {'SAMPLE', 'ID', 'Cage.No'} <= set(clean):
            header = clean
            break
    require(header is not None, '无法定位S2表头')
    selected = []
    for vals in iterator:
        record = dict(zip(header, vals))
        if record.get('SAMPLE') is not None:
            selected.append(record)
    book.close()
    table = pd.DataFrame(selected)
    needed = ['SAMPLE', 'ID', 'Cage.No', 'EarNotch.No', 'Sex', 'Age', 'Genotype', 'Slide.1']
    require(set(needed) <= set(table), 'S2表缺少身份字段：' + str(list(table)))
    table = table[needed].copy()
    table.SAMPLE = table.SAMPLE.map(scalar_id)
    require(not table.SAMPLE.duplicated().any(), 'S2同一样本出现多行，须核查而非取第一行')
    data = source.loc[source.eligible, ['SAMPLE', 'sex', 'age', 'genotype', 'source_file']].merge(
        table, on='SAMPLE', how='left', validate='one_to_one')
    require(data[needed].notna().all().all(), '合格样本缺少发表的动物/笼位身份')
    require((data.sex == data.Sex.astype(str).str.upper().str.strip()).all(), '两来源性别冲突')
    require(np.allclose(data.age, data.Age.map(age_month)), '两来源年龄冲突')
    require((data.genotype == data.Genotype.map(genotype)).all(), '两来源基因型冲突')
    data['animal_id'] = data.ID.astype(str).str.strip()
    data['cage'] = data['Cage.No'].map(cage_id)
    data['slide'] = data['Slide.1'].astype(str).str.strip()
    require(not data.animal_id.duplicated().any(), '同一动物有多个SAMPLE，需要按动物重新定义聚合')
    require(not data[['cage', 'EarNotch.No']].duplicated().any(), '笼号/耳号组合不唯一')
    data[['SAMPLE', 'animal_id', 'Cage.No', 'cage', 'EarNotch.No']].to_csv(
        prepared / 'mouse_cage_labels.csv', index=False)
    print(f'小鼠身份表：{len(data)}个样本，{data.cage.nunique()}个笼位；笼号按原标签保留', flush=True)
    data.to_csv(prepared / 'mouse_manifest.csv', index=False)


def unwrap_mouse(raw):
    """去掉多余外层gzip；给R保留一层压缩，避免写出约20GB的未压缩对象。"""
    path = download(URLS['mouse.rds.gz'], raw / 'mouse.rds.gz')
    layers = []
    for depth in range(4):
        with path.open('rb') as f:
            header = f.read(4)
        if header[:2] != b'\x1f\x8b':
            require(header[:2] in (b'X\n', b'A\n', b'B\n'), '不是R序列化对象')
            break
        with gzip.open(path, 'rb') as f:
            inner = f.read(4)
        layers.append(dict(depth=depth, path=path.name, header=inner.hex()))
        if inner[:2] in (b'X\n', b'A\n', b'B\n'):
            break
        require(inner[:2] == b'\x1f\x8b', '压缩层内部不是gzip或R对象')
        require(shutil.disk_usage(raw).free > 12 * 1024**3, '展开R对象需要至少12GiB可用空间')
        target = raw / f'mouse_layer_{depth+1}.gz'
        temp = target.with_suffix('.tmp')
        with gzip.open(path, 'rb') as src, temp.open('wb') as dst:
            shutil.copyfileobj(src, dst, 8 * 1024**2)
        temp.replace(target)
        path = target
    else:
        raise ValueError('R对象压缩层数异常')
    dump(raw / 'mouse_serialization.json', layers)
    return path


def quality_join(raw, prepared):
    """把官方供体质量信息连接到新聚合的计数单位，交叉检查临床状态。"""
    donors = pd.read_excel(raw / 'seaad_donors.xlsx')
    cols = ['Donor ID', 'Cognitive Status', 'PMI', 'RIN', 'Brain pH', 'Age at Death', 'Sex', 'Braak']
    require(set(cols) <= set(donors), 'SEA-AD供体工作簿列名改变')
    donors = donors[cols].rename(columns={'Donor ID': 'donor_id', 'Brain pH': 'pH'})
    require(not donors.donor_id.duplicated().any(), '供体工作簿编号重复')
    data = pd.read_csv(prepared / 'human' / 'metadata.csv').merge(donors, on='donor_id', how='left', validate='many_to_one')
    a = data.cognitive_status_source.astype(str).str.lower().str.strip()
    b = data['Cognitive Status'].astype(str).str.lower().str.strip()
    require((a == b).all(), 'H5AD与公开供体表临床状态不一致')
    from .seaad import braak
    require(np.allclose(data.age, pd.to_numeric(data['Age at Death'], errors='coerce')), '两份官方资料年龄不一致')
    require((data.sex.astype(str).str.lower() == data.Sex.astype(str).str.lower()).all(), '两份官方资料性别不一致')
    require(np.allclose(data.pathology, data.Braak.map(braak).to_numpy(dtype=float)), '两份官方资料Braak不一致')
    data['clinical_group'] = b.map({'dementia': 'Dementia', 'no dementia': 'No_dementia'})
    for col in ('PMI', 'RIN', 'pH'):
        data[col] = pd.to_numeric(data[col], errors='coerce')
    data.to_csv(prepared / 'human' / 'metadata.csv', index=False)


def read_geo_matrix(path):
    """逐行解析GEO样本特征与表达表；完整读取gzip尾部以检出截断下载。"""
    chars, ids, values, probes, header = [], None, [], [], None
    inside = False
    with gzip.open(path, 'rt', encoding='utf-8') as handle:
        for line in handle:
            parts = next(csv.reader([line], delimiter='\t'))
            if not parts:
                continue
            if parts[0] == '!Sample_geo_accession':
                ids = parts[1:]
            elif parts[0] == '!Sample_characteristics_ch1':
                chars.append(parts[1:])
            elif parts[0] == '!series_matrix_table_begin':
                inside = True
            elif parts[0] == '!series_matrix_table_end':
                inside = False
            elif inside:
                if header is None:
                    header = parts[1:]
                else:
                    probes.append(parts[0])
                    values.append([float(v) for v in parts[1:]])
    require(ids == header and len(set(ids)) == len(ids), 'GEO样本列对应错误')
    metadata = []
    for i, sample in enumerate(ids):
        row = {'sample': sample}
        for char in chars:
            key, sep, value = char[i].partition(':')
            if sep:
                row[key.strip().lower()] = value.strip()
        metadata.append(row)
    return pd.DataFrame(values, index=probes, columns=ids), pd.DataFrame(metadata)


def prepare_panel(raw, prepared):
    """从GEO芯片表达重新构建供体×脑区矩阵，基因面板由R重新拟合产生。"""
    expression, meta = read_geo_matrix(raw / 'GSE118553_series_matrix.txt.gz')
    require({'gender', 'age', 'individual', 'tissue', 'disease state'} <= set(meta), 'GEO样本特征不全')
    with gzip.open(raw / 'GPL10558.annot.gz', 'rt') as handle:
        lines = [line for line in handle if not line.startswith(('#', '!', '^'))]
    annot = pd.read_csv(io.StringIO(''.join(lines)), sep='\t', low_memory=False)
    require({'ID', 'Gene symbol'} <= set(annot), '平台注释列改变：' + str(list(annot)))
    annot = annot[['ID', 'Gene symbol']].dropna()
    annot['gene'] = annot['Gene symbol'].astype(str).str.upper().str.strip()
    annot = annot.loc[annot.gene.str.fullmatch(r'[A-Z0-9][A-Z0-9_.-]*')]
    require(not annot.ID.duplicated().any(), '平台同一探针存在多个注释')
    mapping = annot.set_index('ID').gene
    common = expression.index.intersection(mapping.index)
    require(len(common) / len(expression) > .3, '有效探针映射不足30%')
    expression = expression.loc[common]
    meta['region'] = meta.tissue.str.lower().str.replace('_', ' ')
    meta['group'] = meta['disease state'].str.lower()
    meta['sex'] = meta.gender.str.lower().str.strip()
    meta['age'] = pd.to_numeric(meta.age, errors='coerce')
    meta['donor_id'] = meta.individual.str.replace(r'\s+', ' ', regex=True).str.strip().str.upper()
    meta = meta.loc[meta.region.isin(['frontal cortex', 'entorhinal cortex', 'temporal cortex'])]
    matrices, units = [], []
    for (donor, region), group in meta.groupby(['donor_id', 'region'], sort=True):
        for col in ['group', 'sex', 'age']:
            require(group[col].nunique(dropna=False) == 1, f'技术重复{donor}/{region}的{col}不一致')
        unit = donor + '__' + region.replace(' ', '_')
        probe_values = expression[group['sample']].median(axis=1)
        gene_values = probe_values.groupby(mapping.loc[common].to_numpy()).median()
        gene_values.name = unit
        matrices.append(gene_values)
        units.append(dict(unit_id=unit, donor_id=donor, region=region, group=group.group.iloc[0],
                          sex=group.sex.iloc[0], age=group.age.iloc[0], n_arrays=len(group)))
    matrix = pd.concat(matrices, axis=1)
    require(np.isfinite(matrix.to_numpy()).all(), '芯片表达含缺失值')
    require(np.quantile(matrix.to_numpy(), .99) < 100, '芯片数值不是已处理log表达尺度，须核查GEO处理说明')
    matrix.index.name = 'gene_symbol'
    matrix.to_csv(prepared / 'discovery_expression.csv.gz')
    pd.DataFrame(units).to_csv(prepared / 'discovery_metadata.csv', index=False)
    annot.to_csv(prepared / 'probe_mapping.csv', index=False)


def map_orthologs(raw, results):
    """按MGI同源类别仅保留人鼠一对一符号；无映射端点仍保留在计划中。"""
    source = pd.read_csv(raw / 'orthologs.tsv', sep='\t')
    required = {'DB Class Key', 'Common Organism Name', 'Symbol'}
    require(required <= set(source), 'MGI同源报告格式改变')
    rows = []
    for _, group in source.groupby('DB Class Key'):
        h = group.loc[group['Common Organism Name'] == 'human', 'Symbol'].dropna().unique()
        m = group.loc[group['Common Organism Name'] == 'mouse, laboratory', 'Symbol'].dropna().unique()
        if len(h) == len(m) == 1:
            rows.append((str(h[0]).upper(), str(m[0])))
    mapping = pd.DataFrame(rows, columns=['gene_symbol', 'mouse_gene'])
    mapping = mapping.drop_duplicates()
    mapping = mapping.loc[~mapping.gene_symbol.duplicated(False) & ~mapping.mouse_gene.duplicated(False)]
    panel = pd.read_csv(results / 'panel.csv')
    panel.merge(mapping, how='left', on='gene_symbol', validate='one_to_one').to_csv(results / 'ortholog_panel.csv', index=False)

